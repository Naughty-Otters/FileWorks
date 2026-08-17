"""DOCX and raster image exporters (Markdown / source → download bytes)."""

from __future__ import annotations

import base64
import io
import logging
import re
from html.parser import HTMLParser
from pathlib import Path

from app.processing.export.base import FormatExporter
from app.processing.export.formats import strip_markdown
from app.processing.models import ExportArtifact, OutputKind, ReadResult

logger = logging.getLogger(__name__)

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
_DATA_IMAGE = re.compile(
    r'!\[([^\]]*)\]\((data:image/(png|jpeg|jpg|gif|webp);base64,([A-Za-z0-9+/=\s]+?))'
    r'(?:\s+"([^"]*)")?\)',
    re.IGNORECASE,
)
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_UL = re.compile(r"^[-*+]\s+(.*)$")
_OL = re.compile(r"^(\d+)[.)]\s+(.*)$")
_FENCE = re.compile(r"^```")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _decode_data_image(b64: str) -> bytes | None:
    try:
        return base64.b64decode(re.sub(r"\s+", "", b64), validate=False)
    except Exception:  # noqa: BLE001
        return None


def _docx_image_width(title: str | None) -> float:
    """Convert the editor's pixel-width image title (e.g. ``136 right``) to inches."""
    match = re.match(r"\s*(\d+)", title or "")
    if not match:
        return 5.5
    return max(0.5, min(5.5, int(match.group(1)) / 96))


def markdown_to_docx(markdown: str, *, title: str = "document") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    try:
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    except Exception:  # noqa: BLE001
        pass

    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    i = 0
    in_fence = False
    code_lines: list[str] = []

    def flush_code() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        para = doc.add_paragraph()
        run = para.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        para.paragraph_format.space_after = Pt(10)
        code_lines = []

    def add_styled(paragraph, text: str) -> None:
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            else:
                paragraph.add_run(part)

    while i < len(lines):
        line = lines[i]
        if _FENCE.match(line.strip()):
            if in_fence:
                flush_code()
                in_fence = False
            else:
                in_fence = True
            i += 1
            continue
        if in_fence:
            code_lines.append(line)
            i += 1
            continue

        heading = _HEADING.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            para = doc.add_heading(heading.group(2).strip(), level=level)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and re.search(r"\|?\s*-{3,}", lines[i + 1]):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if re.match(r"^\|?\s*[-:| ]+\|?$", raw):
                    i += 1
                    continue
                cells = [c.strip() for c in raw.strip("|").split("|")]
                rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r_i, row in enumerate(rows):
                    for c_i in range(cols):
                        table.rows[r_i].cells[c_i].text = row[c_i] if c_i < len(row) else ""
                doc.add_paragraph()
            continue

        img = _DATA_IMAGE.search(line)
        if img:
            before = line[: img.start()].strip()
            after = line[img.end() :].strip()
            if before:
                add_styled(doc.add_paragraph(), before)
            payload = _decode_data_image(img.group(4))
            if payload:
                stream = io.BytesIO(payload)
                try:
                    doc.add_picture(stream, width=Inches(_docx_image_width(img.group(5))))
                except Exception:  # noqa: BLE001
                    doc.add_paragraph(img.group(1) or "[image]")
            if after:
                add_styled(doc.add_paragraph(), after)
            i += 1
            continue

        ul = _UL.match(line)
        if ul:
            para = doc.add_paragraph(style="List Bullet")
            add_styled(para, ul.group(1))
            i += 1
            continue
        ol = _OL.match(line)
        if ol:
            para = doc.add_paragraph(style="List Number")
            add_styled(para, ol.group(2))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue
        para = doc.add_paragraph()
        add_styled(para, line.strip())
        i += 1

    flush_code()
    if not doc.paragraphs:
        doc.add_paragraph(title)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _load_font(size: int):
    from PIL import ImageFont

    candidates = (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    )
    for path in candidates:
        if Path(path).is_file():
            try:
                return ImageFont.truetype(path, size=size)
            except OSError:
                continue
    return ImageFont.load_default()


def _stack_images(images: list, *, gap: int = 16, bg=(255, 255, 255)):
    from PIL import Image

    if not images:
        raise ValueError("No images to stack")
    width = max(im.width for im in images)
    height = sum(im.height for im in images) + gap * (len(images) - 1)
    canvas = Image.new("RGB", (width, height), bg)
    y = 0
    for im in images:
        if im.mode != "RGB":
            im = im.convert("RGB")
        x = (width - im.width) // 2
        canvas.paste(im, (x, y))
        y += im.height + gap
    return canvas


def _rasterize_pdf(source_bytes: bytes, *, max_pages: int = 12, scale: float = 1.6):
    import pypdfium2 as pdfium
    from PIL import Image

    document = pdfium.PdfDocument(source_bytes)
    pages: list[Image.Image] = []
    try:
        count = min(len(document), max_pages)
        for index in range(count):
            page = document[index]
            bitmap = page.render(scale=scale)
            pages.append(bitmap.to_pil().convert("RGB"))
            page.close()
    finally:
        document.close()
    return pages


def _open_source_image(source_bytes: bytes):
    from PIL import Image

    image = Image.open(io.BytesIO(source_bytes))
    image.load()
    return image.convert("RGB")


def _wrap_text(text: str, font, max_width: int, draw) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        words = raw.split(" ")
        current = ""
        for word in words:
            trial = word if not current else f"{current} {word}"
            if draw.textlength(trial, font=font) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        lines.append(current)
    return lines or [""]


def _render_markdown_image(markdown: str) -> object:
    from PIL import Image, ImageDraw

    width = 960
    padding = 48
    font = _load_font(18)
    heading_font = _load_font(28)
    text = strip_markdown(markdown or "") or "Converted document"
    probe = Image.new("RGB", (width, 100), (255, 255, 255))
    draw = ImageDraw.Draw(probe)
    wrapped = _wrap_text(text, font, width - padding * 2, draw)
    line_h = 26
    extra: list[Image.Image] = []
    for match in _DATA_IMAGE.finditer(markdown or ""):
        payload = _decode_data_image(match.group(4))
        if not payload:
            continue
        try:
            im = Image.open(io.BytesIO(payload)).convert("RGB")
            max_w = width - padding * 2
            if im.width > max_w:
                ratio = max_w / im.width
                im = im.resize((max_w, max(1, int(im.height * ratio))), Image.Resampling.LANCZOS)
            extra.append(im)
        except Exception:  # noqa: BLE001
            continue
        if len(extra) >= 4:
            break

    text_h = max(120, padding * 2 + len(wrapped) * line_h + 40)
    img_h = sum(im.height + 16 for im in extra)
    canvas = Image.new("RGB", (width, text_h + img_h), (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    draw.text((padding, padding), wrapped[0][:80] if wrapped else "Document", font=heading_font, fill=(26, 26, 26))
    y = padding + 44
    for line in wrapped[:80]:
        draw.text((padding, y), line, font=font, fill=(40, 40, 40))
        y += line_h
    for im in extra:
        canvas.paste(im, (padding, y))
        y += im.height + 16
    return canvas


def render_document_image(
    markdown: str,
    *,
    image_format: str,
    source_bytes: bytes | None = None,
    source_suffix: str | None = None,
) -> bytes:
    from PIL import Image

    suffix = (source_suffix or "").lower()
    image: Image.Image | None = None
    if source_bytes and suffix == ".pdf":
        try:
            pages = _rasterize_pdf(source_bytes)
            if pages:
                image = _stack_images(pages)
        except Exception:  # noqa: BLE001
            logger.exception("PDF rasterize failed; falling back to Markdown preview")
    elif source_bytes and suffix in _IMAGE_SUFFIXES:
        try:
            image = _open_source_image(source_bytes)
        except Exception:  # noqa: BLE001
            logger.exception("Image convert failed; falling back to Markdown preview")
    if image is None:
        image = _render_markdown_image(markdown)

    max_side = 2400
    w, h = image.size
    longest = max(w, h)
    if longest > max_side:
        ratio = max_side / longest
        image = image.resize((max(1, int(w * ratio)), max(1, int(h * ratio))), Image.Resampling.LANCZOS)

    buffer = io.BytesIO()
    fmt = "JPEG" if image_format.lower() in {"jpeg", "jpg"} else "PNG"
    if fmt == "JPEG":
        image.convert("RGB").save(buffer, format="JPEG", quality=85, optimize=True)
    else:
        image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


class _DocxRichTextParser(HTMLParser):
    """Small safe inline-HTML → DOCX run mapper for OCR formatting."""

    def __init__(self, paragraph) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraph = paragraph
        self.stack: list[dict[str, object]] = [{}]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.lower()
        if name == "br":
            self.paragraph.add_run().add_break()
            return
        state = dict(self.stack[-1])
        attr = {key.lower(): (value or "") for key, value in attrs}
        style = attr.get("style", "").lower()
        css = {
            key.lower(): value.strip()
            for key, value in re.findall(r"([a-zA-Z-]+)\s*:\s*([^;]+)", style)
        }
        if name in {"b", "strong"} or "font-weight:bold" in style or "font-weight: bold" in style:
            state["bold"] = True
        if name in {"i", "em"} or "font-style:italic" in style or "font-style: italic" in style:
            state["italic"] = True
        if name == "u" or "text-decoration:underline" in style or "text-decoration: underline" in style:
            state["underline"] = True
        if name == "mark" or "background" in style:
            state["highlight"] = True
        if name in {"s", "strike", "del"} or "line-through" in css.get("text-decoration", ""):
            state["strike"] = True
        if name == "sup" or css.get("vertical-align") in {"super", "sup"}:
            state["superscript"] = True
            state["subscript"] = False
        if name == "sub" or css.get("vertical-align") == "sub":
            state["subscript"] = True
            state["superscript"] = False
        size = re.search(r"font-size\s*:\s*(\d+(?:\.\d+)?)\s*(px|pt)?", style)
        if size:
            value = float(size.group(1))
            state["size_pt"] = value * 0.75 if size.group(2) != "pt" else value
        if css.get("color"):
            state["color"] = css["color"]
        if css.get("font-family"):
            state["font_family"] = css["font-family"].strip("'\"").split(",")[0].strip()
        self.stack.append(state)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() != "br" and len(self.stack) > 1:
            self.stack.pop()

    def handle_data(self, data: str) -> None:
        if not data:
            return
        from docx.enum.text import WD_COLOR_INDEX
        from docx.shared import Pt, RGBColor

        run = self.paragraph.add_run(data)
        state = self.stack[-1]
        run.bold = bool(state.get("bold"))
        run.italic = bool(state.get("italic"))
        run.underline = bool(state.get("underline"))
        run.font.strike = bool(state.get("strike"))
        run.font.superscript = bool(state.get("superscript"))
        run.font.subscript = bool(state.get("subscript"))
        if state.get("highlight"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if isinstance(state.get("size_pt"), (int, float)):
            run.font.size = Pt(float(state["size_pt"]))
        color = _parse_rgb(str(state.get("color") or ""))
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
        if state.get("font_family"):
            run.font.name = str(state["font_family"])


def _parse_rgb(value: str) -> tuple[int, int, int] | None:
    text = value.strip().lower()
    named = {
        "black": (0, 0, 0),
        "white": (255, 255, 255),
        "red": (255, 0, 0),
        "green": (0, 128, 0),
        "blue": (0, 0, 255),
        "yellow": (255, 255, 0),
        "gray": (128, 128, 128),
        "grey": (128, 128, 128),
    }
    if text in named:
        return named[text]
    match = re.fullmatch(r"#([0-9a-f]{3}|[0-9a-f]{6})", text)
    if match:
        raw = match.group(1)
        if len(raw) == 3:
            raw = "".join(char * 2 for char in raw)
        return int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    match = re.fullmatch(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", text)
    if match:
        return tuple(max(0, min(255, int(part))) for part in match.groups())  # type: ignore[return-value]
    return None


def _add_rich_html(paragraph, rich: str | None, fallback: str = "") -> None:
    if rich:
        parser = _DocxRichTextParser(paragraph)
        try:
            parser.feed(rich)
            return
        except Exception:  # noqa: BLE001
            pass
    paragraph.add_run(fallback)


def structure_to_docx(structure, *, title: str = "document", markdown_fallback: str = "") -> bytes:
    """Build DOCX from StructureDocument when blocks carry tables/headings."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    from app.processing.structure.schema import StructureDocument

    if not isinstance(structure, StructureDocument):
        return markdown_to_docx(markdown_fallback or "", title=title)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    try:
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    except Exception:  # noqa: BLE001
        pass

    title_map = {
        "doc_title": 1,
        "paragraph_title": 2,
        "figure_title": 3,
        "table_title": 3,
        "chart_title": 3,
        "content_title": 2,
    }
    ignore = {
        "number",
        "header",
        "header_image",
        "footer",
        "footer_image",
        "formula_number",
    }
    blocks = sorted(
        structure.blocks,
        key=lambda b: (
            b.page,
            b.order is None,
            b.order if isinstance(b.order, (int, float)) else 10**9,
        ),
    )
    wrote = False

    def apply_alignment(paragraph, alignment: str | None) -> None:
        value = (alignment or "").lower()
        mapping = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        if value in mapping:
            paragraph.alignment = mapping[value]

    for block in blocks:
        label = (block.label or "text").lower()
        if label in ignore:
            continue
        text = (block.text or "").strip()
        if label in title_map and text:
            para = doc.add_heading("", level=min(block.title_level or title_map[label], 4))
            _add_rich_html(para, block.rich_text, text.lstrip("#").strip())
            apply_alignment(para, block.alignment)
            wrote = True
            continue
        if label == "table" and block.table and block.table.cells:
            table = block.table
            rows = max(table.rows, max((c.r + c.rs for c in table.cells), default=0))
            cols = max(table.cols, max((c.c + c.cs for c in table.cells), default=0))
            if rows > 0 and cols > 0:
                grid = doc.add_table(rows=rows, cols=cols)
                grid.style = "Table Grid"
                for cell in table.cells:
                    try:
                        target = grid.rows[cell.r].cells[cell.c]
                        target.text = ""
                        _add_rich_html(target.paragraphs[0], cell.html, cell.text or "")
                        apply_alignment(target.paragraphs[0], cell.align)
                        if cell.is_header:
                            for run in target.paragraphs[0].runs:
                                run.bold = True
                        if cell.rs > 1 or cell.cs > 1:
                            target.merge(
                                grid.rows[min(rows - 1, cell.r + cell.rs - 1)].cells[
                                    min(cols - 1, cell.c + cell.cs - 1)
                                ]
                            )
                    except Exception:  # noqa: BLE001
                        continue
                doc.add_paragraph()
                wrote = True
            continue
        if label == "formula" and (block.formula or text):
            para = doc.add_paragraph(block.formula or text)
            para.paragraph_format.space_after = Pt(10)
            wrote = True
            continue
        if label in {"image", "chart", "seal"}:
            ref = block.image_ref or text
            img = _DATA_IMAGE.search(ref)
            if img:
                payload = _decode_data_image(img.group(4))
                if payload:
                    stream = io.BytesIO(payload)
                    try:
                        doc.add_picture(stream, width=Inches(_docx_image_width(img.group(5))))
                        wrote = True
                        continue
                    except Exception:  # noqa: BLE001
                        pass
            if text:
                doc.add_paragraph(text)
                wrote = True
            continue
        if text:
            # Fall back to markdown line handling for mixed content.
            if text.startswith("#") or "|" in text or text.startswith("!["):
                # Append via markdown helper by merging — simpler: plain paragraph.
                for line in text.splitlines():
                    if line.strip():
                        doc.add_paragraph(line.strip())
                wrote = True
            else:
                para = doc.add_paragraph()
                _add_rich_html(para, block.rich_text, text)
                apply_alignment(para, block.alignment)
                wrote = True

    if not wrote:
        return markdown_to_docx(markdown_fallback or project_md(structure), title=title)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def project_md(structure) -> str:
    from app.processing.structure.project import project_markdown

    return project_markdown(structure)


class DocxExporter(FormatExporter):
    id = "docx"
    label = "Word"
    aliases = ("word",)

    def can_export(self, source_kind: OutputKind = OutputKind.MARKDOWN) -> bool:
        return source_kind in {OutputKind.MARKDOWN, OutputKind.TEXT}

    def export(
        self,
        *,
        markdown: str,
        title: str = "document",
        read_result: ReadResult | None = None,
        source_bytes: bytes | None = None,
        source_suffix: str | None = None,
        structure=None,
    ) -> ExportArtifact:
        del read_result, source_bytes, source_suffix
        from app.processing.structure.richness import structure_is_richer

        data = (
            structure_to_docx(structure, title=title, markdown_fallback=markdown or "")
            if structure_is_richer(structure, markdown)
            else markdown_to_docx(markdown or "", title=title)
        )
        return ExportArtifact(
            data=data,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            extension="docx",
            format_id=self.id,
        )


class PngExporter(FormatExporter):
    id = "png"
    label = "PNG image"
    aliases = ()

    def can_export(self, source_kind: OutputKind = OutputKind.MARKDOWN) -> bool:
        return True

    def export(
        self,
        *,
        markdown: str,
        title: str = "document",
        read_result: ReadResult | None = None,
        source_bytes: bytes | None = None,
        source_suffix: str | None = None,
        structure=None,
    ) -> ExportArtifact:
        del title, read_result, structure
        return ExportArtifact(
            data=render_document_image(
                markdown or "",
                image_format="png",
                source_bytes=source_bytes,
                source_suffix=source_suffix,
            ),
            media_type="image/png",
            extension="png",
            format_id=self.id,
        )


class JpegExporter(FormatExporter):
    id = "jpeg"
    label = "JPEG image"
    aliases = ("jpg",)

    def can_export(self, source_kind: OutputKind = OutputKind.MARKDOWN) -> bool:
        return True

    def export(
        self,
        *,
        markdown: str,
        title: str = "document",
        read_result: ReadResult | None = None,
        source_bytes: bytes | None = None,
        source_suffix: str | None = None,
        structure=None,
    ) -> ExportArtifact:
        del title, read_result, structure
        return ExportArtifact(
            data=render_document_image(
                markdown or "",
                image_format="jpeg",
                source_bytes=source_bytes,
                source_suffix=source_suffix,
            ),
            media_type="image/jpeg",
            extension="jpg",
            format_id=self.id,
        )
