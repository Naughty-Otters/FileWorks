from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any

from app.processing.template.schema import DesignDocument, DesignNode


def _style(node: DesignNode) -> str:
    style = node.style
    pairs: list[tuple[str, Any]] = [
        ("position", "absolute"),
        ("left", f"{node.x}px"),
        ("top", f"{node.y}px"),
        ("width", f"{node.w}px"),
        ("height", f"{node.h}px"),
        ("font-family", style.font_family),
        ("font-size", f"{style.font_size}px" if style.font_size else None),
        ("font-weight", "bold" if style.bold else None),
        ("font-style", "italic" if style.italic else None),
        ("text-decoration", "underline" if style.underline else None),
        ("color", style.color),
        ("background", style.background_color),
        ("text-align", style.align),
        (
            "border",
            f"{style.border_width}px {style.border_style} {style.border_color}",
        ),
        ("object-fit", node.object_fit if node.type == "image" else None),
        ("box-sizing", "border-box"),
        ("overflow-wrap", "anywhere"),
    ]
    return ";".join(f"{key}:{value}" for key, value in pairs if value is not None)


def _html_node(node: DesignNode) -> str:
    attrs = f'id="{escape(node.id, quote=True)}" style="{escape(_style(node), quote=True)}"'
    if node.type == "image":
        if not node.image_src:
            return ""
        src = node.image_src.strip()
        if not src.startswith(("https://", "http://", "data:image/", "/")):
            return ""
        return (
            f'<img {attrs} src="{escape(src, quote=True)}" '
            f'alt="{escape(node.label, quote=True)}">'
        )
    children = "".join(_html_node(child) for child in node.children)
    text = "" if node.type == "box" else escape(node.label).replace("\n", "<br>")
    if node.type == "heading":
        tag = f"h{node.level}"
    elif node.type == "paragraph":
        tag = "p"
    else:
        tag = "div"
    return f"<{tag} {attrs}>{text}{children}</{tag}>"


def render_html(document: DesignDocument, *, title: str = "Document") -> str:
    page = document.page
    nodes = _html_node(document.root)
    rendered_page = (
        f'<section class="page" data-page="1" style="width:{page.w}px;height:{page.h}px">'
        f"{nodes}</section>"
    )
    return (
        "<!doctype html><html><head><meta charset=\"utf-8\">"
        f"<title>{escape(title)}</title><style>"
        "html,body{margin:0;background:#eee}.page{position:relative;margin:16px auto;"
        "overflow:hidden;box-shadow:0 1px 5px #999}@media print{body{background:white}"
        ".page{margin:0;box-shadow:none;break-after:page}}</style></head><body>"
        + rendered_page
        + "</body></html>"
    )


def render_docx(document: DesignDocument, *, title: str = "Document") -> bytes:
    """Editable, intentionally linear DOCX representation of the design tree."""
    from docx import Document

    output = Document()
    output.core_properties.title = title

    def add(nodes: list[DesignNode], depth: int = 0) -> None:
        for node in nodes:
            if node.type == "heading":
                output.add_heading(node.label, level=min(node.level, 9))
            elif node.type == "paragraph":
                output.add_paragraph(node.label)
            elif node.type == "image":
                alt = getattr(node, "alt", None) or node.label or f"[Image: {node.image_src or ''}]"
                output.add_paragraph(alt)
            add(node.children, depth + 1)

    add([document.root])
    buffer = BytesIO()
    output.save(buffer)
    return buffer.getvalue()


def render_markdown(document: DesignDocument) -> str:
    """Linear Markdown of the expanded design tree (filled content, not design JSON)."""
    lines: list[str] = []

    def walk(node: DesignNode) -> None:
        if node.type == "heading":
            level = max(1, min(int(node.level or 1), 6))
            text = (node.label or "").strip()
            if text:
                lines.append(f"{'#' * level} {text}")
                lines.append("")
        elif node.type == "paragraph":
            text = (node.label or "").strip()
            if text:
                lines.append(text)
                lines.append("")
        elif node.type == "image":
            src = (node.image_src or "").strip()
            alt = (getattr(node, "alt", None) or node.label or "image").strip()
            if src:
                lines.append(f"![{alt}]({src})")
                lines.append("")
        for child in node.children:
            walk(child)

    walk(document.root)
    return "\n".join(lines).strip() + "\n"


def _parse_color(
    value: str | None, default: tuple[float, float, float] | None = (0, 0, 0)
) -> tuple[float, float, float] | None:
    if value is None:
        return default
    if value in {"transparent", "none", ""}:
        return None
    raw = value.strip().lstrip("#")
    if len(raw) == 3:
        raw = "".join(ch * 2 for ch in raw)
    if len(raw) != 6:
        return default
    try:
        return tuple(int(raw[i : i + 2], 16) / 255.0 for i in (0, 2, 4))  # type: ignore[return-value]
    except ValueError:
        return default


def render_pdf(document: DesignDocument, *, title: str = "Document") -> bytes:
    """Absolute-position PDF from the laid-out design tree (ReportLab)."""
    from reportlab.lib.colors import Color
    from reportlab.pdfgen import canvas

    from app.processing.template.layout import TEXT_PAD_X, TEXT_PAD_Y

    buffer = BytesIO()
    page_w = float(document.page.w)
    page_h = float(document.page.h)
    total_h = max(page_h, float(document.root.h or page_h))
    page_count = max(1, int(-(-total_h // page_h)))

    pdf = canvas.Canvas(buffer, pagesize=(page_w, page_h))
    pdf.setTitle(title)

    def draw_node(node: DesignNode, origin_x: float, origin_y: float, page_index: int) -> None:
        abs_x = origin_x + node.x
        abs_y = origin_y + node.y
        page_top = page_index * page_h
        page_bottom = page_top + page_h
        if not (abs_y + node.h < page_top or abs_y > page_bottom):
            local_y = abs_y - page_top
            pdf_x = abs_x
            pdf_y = page_h - local_y - node.h

            bg = _parse_color(
                None if node.style.background_color in {"transparent", "none", ""} else node.style.background_color,
                default=(1, 1, 1),
            )
            if node.style.background_color not in {"transparent", "none", "", None} and bg is not None:
                pdf.setFillColor(Color(*bg))
                pdf.rect(pdf_x, pdf_y, node.w, node.h, stroke=0, fill=1)

            border_w = float(node.style.border_width or 0)
            border_c = _parse_color(node.style.border_color)
            if border_w > 0 and border_c is not None and node.style.border_style != "none":
                pdf.setStrokeColor(Color(*border_c))
                pdf.setLineWidth(border_w)
                if node.style.border_style == "dashed":
                    pdf.setDash(3, 2)
                else:
                    pdf.setDash()
                pdf.rect(pdf_x, pdf_y, node.w, node.h, stroke=1, fill=0)
                pdf.setDash()

            if node.type == "image" and node.image_src:
                src = node.image_src.strip()
                try:
                    if src.startswith("data:image/"):
                        import base64
                        import re

                        match = re.match(r"data:image/[^;]+;base64,(.+)", src, re.S)
                        if match:
                            from reportlab.lib.utils import ImageReader

                            img = ImageReader(BytesIO(base64.b64decode(match.group(1))))
                            pdf.drawImage(
                                img,
                                pdf_x,
                                pdf_y,
                                width=node.w,
                                height=node.h,
                                preserveAspectRatio=node.object_fit == "contain",
                                mask="auto",
                            )
                except Exception:
                    pass
            elif node.type in {"heading", "paragraph"} and (node.label or "").strip():
                color = _parse_color(node.style.color) or (0, 0, 0)
                pdf.setFillColor(Color(*color))
                font_size = float(node.style.font_size or (18 if node.type == "heading" else 12))
                font_name = "Helvetica-Bold" if node.style.bold else "Helvetica"
                if node.style.italic:
                    font_name = "Helvetica-Oblique" if not node.style.bold else "Helvetica-BoldOblique"
                pdf.setFont(font_name, font_size)
                text = node.label.replace("\r\n", "\n").replace("\r", "\n")
                lines = text.split("\n")
                line_gap = font_size * (1.25 if node.type == "heading" else 1.4)
                cursor = pdf_y + node.h - TEXT_PAD_Y - font_size
                max_width = max(10.0, node.w - 2 * TEXT_PAD_X)
                for line in lines:
                    avg = font_size * 0.52
                    capacity = max(1, int(max_width / avg))
                    chunks = [line[i : i + capacity] for i in range(0, max(1, len(line)), capacity)] or [""]
                    for chunk in chunks:
                        if cursor < pdf_y:
                            break
                        if node.style.align == "center":
                            pdf.drawCentredString(pdf_x + node.w / 2, cursor, chunk)
                        elif node.style.align == "right":
                            pdf.drawRightString(pdf_x + node.w - TEXT_PAD_X, cursor, chunk)
                        else:
                            pdf.drawString(pdf_x + TEXT_PAD_X, cursor, chunk)
                        cursor -= line_gap

        child_origin_x = 0.0 if node.id == "root" else abs_x
        child_origin_y = 0.0 if node.id == "root" else abs_y
        for child in node.children:
            draw_node(child, child_origin_x, child_origin_y, page_index)

    for page_index in range(page_count):
        if page_index > 0:
            pdf.showPage()
        draw_node(document.root, 0.0, 0.0, page_index)

    pdf.save()
    return buffer.getvalue()
